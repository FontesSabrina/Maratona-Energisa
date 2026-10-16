"""Teste do portão e da leitura segura do arquivo enviado (api/leitura.py), sem banco.

Arquivos maliciosos gerados por script, cada um recusado com a mensagem fixa certa (sem conteúdo do
arquivo) ou lido sem efeito colateral:
- zip bomb (descompactado acima de 50 MB; razão acima de 100), arquivos internos demais, caminhos com
  ".." ou absolutos, zip cifrado, índice mentiroso, macro (.xlsm: vbaProject.bin ou macroEnabled),
  .xlsb, zip que não é planilha, .exe renomeado, .xls/xlsx protegido por senha (OLE), PDF;
- XML com entidades (billion laughs / XXE) barrado pelo defusedxml;
- fórmula sem valor guardado (marcada para a linha ser rejeitada depois), célula com "<script>",
  célula gigante (numa coluna mantida chega como está; numa descartada nem sai da leitura);
- .csv com bytes nulos, codificação inválida, Windows-1252, UTF-8 com BOM, separador vírgula,
  "=HYPERLINK(...)" lido como texto;
- aba oculta (e muito oculta) com dados nunca aparece; uma aba visível que serve é escolhida sozinha;
  com duas, a lista volta para a pessoa escolher; aba escolhida inexistente é recusada;
- limites: mais de 60 colunas, mais de 50.000 linhas, tempo máximo;
- minimização: da aba Postes do exemplo só saem as 15 colunas do modelo; NOME, NUM_CASA, MEDIDOR,
  NUM_CONTA... não saem da leitura (nem o nome da coluna, nem os valores);
- um .xlsx de 50.000 linhas passa.

Uso (na raiz): uv run python -W ignore testes/api/teste_leitura.py
"""

import io
import re
import struct
import sys
import time
import zipfile
from pathlib import Path

import openpyxl
import openpyxl.xml

from comunidades.api import leitura
from comunidades.nucleo.motivos import FORMULA_SEM_VALOR
from comunidades.nucleo.planilha import PlanilhaInvalida

resultados = []
EXEMPLOS = Path(__file__).resolve().parents[1] / "planilha" / "exemplos"


def conferir(nome, ok, extra=""):
    resultados.append(("OK  " if ok else "FALHA", nome, extra))


def xlsx(abas: dict[str, list[list]], estados: dict[str, str] | None = None) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nome, linhas in abas.items():
        ws = wb.create_sheet(nome)
        for linha in linhas:
            ws.append(linha)
        ws.sheet_state = (estados or {}).get(nome, "visible")
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def refazer_zip(base: bytes, extra: dict[str, bytes] | None = None, trocar: dict[str, bytes] | None = None) -> bytes:
    """Copia um .xlsx trocando ou acrescentando arquivos internos."""
    b = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(base)) as z, zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as w:
        for i in z.infolist():
            w.writestr(i.filename, (trocar or {}).get(i.filename, z.read(i.filename)))
        for nome, dados in (extra or {}).items():
            w.writestr(nome, dados)
    return b.getvalue()


def zip_com(itens: dict[str, bytes]) -> bytes:
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as z:
        for nome, dados in itens.items():
            z.writestr(nome, dados)
    return b.getvalue()


def recusa(conteudo: bytes, aba=None) -> str | None:
    """O código da recusa (None: aceito)."""
    try:
        leitura.ler_tabela(conteudo, aba)
    except leitura.Recusa as e:
        assert str(e) in leitura.MENSAGENS.values()
        return e.codigo
    return None


CAB = ["uc", "latitude", "longitude", "equipamento_vital", "classe", "local"]
LINHA = ["0000000001", -21.530443, -42.638946, "não", "residencial", "Rua das Flores"]
NORMAL = xlsx({"Modelo": [CAB, LINHA]})
t = leitura.ler_tabela(NORMAL)
conferir(".xlsx normal: lido, aba escolhida sozinha", t.formato == "xlsx" and t.aba == "Modelo" and t.linhas == [(2, LINHA)], repr(t))
conferir("defusedxml ligado no openpyxl", openpyxl.xml.DEFUSEDXML is True)

# ---------- portão ----------
TIPOS = 'xmlns="http://schemas.openxmlformats.org/package/2006/content-types"'
casos = {
    "vazio": b"   \n ",
    "grande": b"uc;latitude;longitude\n" + b"1;-21,5;-42,6\n" * (leitura.MAX_BYTES // 14 + 1),
    "ole": bytes.fromhex("D0CF11E0A1B11AE1") + b"\0" * 600,
    ".exe renomeado": b"MZ\x90\x00\x03\x00\x00\x00" + b"\0" * 200,
    "PDF": b"%PDF-1.7\n" + b"x" * 100,
    "bomba (60 MB de zeros)": zip_com({"[Content_Types].xml": b"<Types/>", "xl/workbook.xml": b"<w/>",
                                       "xl/worksheets/sheet1.xml": b"\0" * (60 * 1024 * 1024)}),
    "razão ~1000 (10 MB de zeros)": zip_com({"[Content_Types].xml": b"<Types/>", "xl/workbook.xml": b"<w/>",
                                             "xl/worksheets/sheet1.xml": b"\0" * (10 * 1024 * 1024)}),
    "201 arquivos internos": zip_com({f"xl/x{i}.xml": b"<a/>" for i in range(leitura.MAX_ENTRADAS + 1)}),
    "caminho com ..": refazer_zip(NORMAL, extra={"xl/../../evil.xml": b"<a/>"}),
    "caminho absoluto": refazer_zip(NORMAL, extra={"/etc/evil.xml": b"<a/>"}),
    "caminho com unidade": refazer_zip(NORMAL, extra={"C:/evil.xml": b"<a/>"}),
    "macro (vbaProject.bin)": refazer_zip(NORMAL, extra={"xl/vbaProject.bin": b"\0\1\2"}),
    "macro (macroEnabled)": refazer_zip(NORMAL, trocar={"[Content_Types].xml": f'<Types {TIPOS}><Override PartName="/xl/workbook.xml" '
                                                       'ContentType="application/vnd.ms-excel.sheet.macroEnabled.main+xml"/></Types>'.encode()}),
    ".xlsb": zip_com({"[Content_Types].xml": b"<Types/>", "xl/workbook.bin": b"\0\1"}),
    "zip que não é planilha": zip_com({"[Content_Types].xml": b"<Types/>", "word/document.xml": b"<d/>"}),
    "zip corrompido": b"PK\x03\x04lixo" + b"\0" * 50,
    "csv com NUL": b"uc;latitude;longitude\n1;-21,5;-42,6\0\n",
    "csv com codificação inválida": "uc;latitude;longitude\n1;-21,5;-42,6\n".encode() + b"\x81\x8d\xff\xfe",
    "csv sem as colunas": b"nome;telefone\nFulano;123\n",
    "61 colunas": xlsx({"A": [CAB + [f"c{i}" for i in range(55)], LINHA]}),
    "nenhuma aba serve": xlsx({"A": [["nome", "rua"], ["x", "y"]]}),
}
esperado = {"vazio": "vazio", "grande": "grande", "ole": "ole", ".exe renomeado": "nao_planilha", "PDF": "nao_planilha",
            "bomba (60 MB de zeros)": "descompactado", "razão ~1000 (10 MB de zeros)": "razao",
            "201 arquivos internos": "entradas", "caminho com ..": "caminho", "caminho absoluto": "caminho",
            "caminho com unidade": "caminho", "macro (vbaProject.bin)": "macro", "macro (macroEnabled)": "macro",
            ".xlsb": "xlsb", "zip que não é planilha": "nao_planilha", "zip corrompido": "zip", "csv com NUL": "nul",
            "csv com codificação inválida": "codificacao", "csv sem as colunas": "sem_colunas", "61 colunas": "colunas",
            "nenhuma aba serve": "sem_aba"}
for nome, conteudo in casos.items():
    r = recusa(conteudo)
    conferir(f"recusado: {nome} ({esperado[nome]})", r == esperado[nome], str(r))

# zip cifrado: liga o bit de cifra no índice central
b = bytearray(NORMAL)
cd = b.find(b"PK\x01\x02")
struct.pack_into("<H", b, cd + 8, struct.unpack_from("<H", b, cd + 8)[0] | 1)
conferir("recusado: zip cifrado (cifrado)", recusa(bytes(b)) == "cifrado", str(recusa(bytes(b))))

# índice mentiroso: declara 1.000 bytes para uma planilha de 5 MB
real = "".join(f"<c r='A{i}'><v>{i * 7919 % 100003}</v></c>" for i in range(200_000)).encode()[:5 * 1024 * 1024]
b = bytearray(refazer_zip(NORMAL, trocar={"xl/worksheets/sheet1.xml": real}))
pos = next(m.start() for m in re.finditer(re.escape(b"PK\x01\x02"), bytes(b))
           if b[m.start() + 46:m.start() + 46 + 24] == b"xl/worksheets/sheet1.xml")   # a entrada no índice central
struct.pack_into("<I", b, pos + 24, 1000)
conferir("recusado: índice mentiroso (o zipfile não passa do declarado)", recusa(bytes(b)) == "zip", str(recusa(bytes(b))))

# billion laughs / XXE no XML da planilha (o openpyxl grava os textos na própria aba): o defusedxml barra
NS = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
def aba_com(dtd: str, entidade: str) -> bytes:
    return (f'<?xml version="1.0"?>{dtd}<worksheet {NS}><sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>uc</t></is></c>'
            f'<c r="B1" t="inlineStr"><is><t>latitude</t></is></c><c r="C1" t="inlineStr"><is><t>longitude</t></is></c></row>'
            f'<row r="2"><c r="A2" t="inlineStr"><is><t>{entidade}</t></is></c><c r="B2"><v>-21.5</v></c><c r="C2"><v>-42.6</v></c>'
            '</row></sheetData></worksheet>').encode()
risada = aba_com('<!DOCTYPE lolz [<!ENTITY lol "lol"><!ENTITY lol2 "&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;">'
                 '<!ENTITY lol3 "&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;">]>', "&lol3;")
xxe = aba_com('<!DOCTYPE x [<!ENTITY xxe SYSTEM "file:///C:/Windows/win.ini">]>', "&xxe;")
base3 = xlsx({"Modelo": [["uc", "latitude", "longitude"], ["texto", -21.5, -42.6]]})
for nome, xml in [("billion laughs", risada), ("XXE", xxe)]:
    r = recusa(refazer_zip(base3, trocar={"xl/worksheets/sheet1.xml": xml}))
    conferir(f"recusado: XML com entidades ({nome}), barrado pelo defusedxml", r == "zip", str(r))
conferir("controle: a mesma aba sem entidades é lida", recusa(refazer_zip(base3, trocar={"xl/worksheets/sheet1.xml": aba_com("", "1")})) is None)

# ---------- conteúdo estranho, lido sem efeito ----------
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Modelo"
ws.append(CAB + ["NOME"])
ws.append(["1", "=1-22.5", -42.63, "não", "<script>alert(1)</script>", "=HYPERLINK(\"http://x\")", "x" * 100_000])
ws.append([None] * 7)
ws.append(["2", -21.53, -42.63, None, None, "Z" * 100_000, "Fulano de Tal"])
b = io.BytesIO()
wb.save(b)
t = leitura.ler_tabela(b.getvalue())
lin = dict(t.linhas)
conferir("fórmula sem valor guardado marcada (FORMULA_SEM_VALOR)", lin[2][t.campos["latitude"]] is FORMULA_SEM_VALOR, repr(lin[2]))
conferir("\"<script>\" chega como texto, sem ser interpretado", lin[2][t.campos["classe"]] == "<script>alert(1)</script>")
conferir("\"=HYPERLINK\" no .xlsx é fórmula sem valor: marcada, nunca avaliada", lin[2][t.campos["local"]] is FORMULA_SEM_VALOR, repr(lin[2]))
conferir("célula gigante numa coluna mantida chega (32.767, o limite do Excel; a linha é rejeitada no tradutor)",
         len(lin[4][t.campos["local"]]) == 32_767, str(len(lin[4][t.campos["local"]])))
conferir("célula gigante numa coluna descartada nem sai da leitura", "NOME" not in t.cabecalho and t.colunas_ignoradas == 1
         and "Fulano de Tal" not in repr(t.linhas) and "x" * 1000 not in repr(t.linhas))
conferir("linha vazia pulada, numeração da planilha mantida (2 e 4)", sorted(lin) == [2, 4], str(sorted(lin)))

# abas ocultas
dados_ocultos = ["SEGREDO-OCULTO", -21.53, -42.63, "sim", "x", "Rua Oculta"]
oculta = xlsx({"Modelo": [CAB, LINHA], "Oculta": [CAB, dados_ocultos], "MuitoOculta": [CAB, dados_ocultos]},
              {"Oculta": "hidden", "MuitoOculta": "veryHidden"})
t = leitura.ler_tabela(oculta)
conferir("abas oculta e muito oculta com dados: ignoradas (sem pedir escolha)", t.aba == "Modelo" and "SEGREDO" not in repr(t), repr(t.aba))
conferir("aba oculta não pode ser escolhida pelo nome", recusa(oculta, "Oculta") == "aba_invalida")
so_oculta = xlsx({"Notas": [["texto"]], "Oculta": [CAB, dados_ocultos]}, {"Oculta": "hidden"})
conferir("só a aba oculta serve: recusado (sem_aba)", recusa(so_oculta) == "sem_aba")

# ---------- exemplos oficiais ----------
postes = (EXEMPLOS / "planilha_postes.xlsx").read_bytes()
try:
    leitura.ler_tabela(postes)
    conferir("exemplo com Modelo e Postes: pede a escolha da aba", False)
except leitura.EscolhaDeAba as e:
    conferir("exemplo com Modelo e Postes: pede a escolha da aba", e.abas == ["Modelo", "Postes"], str(e.abas))
for aba, ruim in [("Notas", "aba_invalida"), ("Inexistente", "aba_invalida")]:
    conferir(f"aba escolhida {aba!r}: recusada", recusa(postes, aba) == ruim)
tm = leitura.ler_tabela(postes, "Modelo")
tp = leitura.ler_tabela(postes, "Postes")
conferir("aba Modelo: 200 linhas, as 6 colunas, nenhuma ignorada", len(tm.linhas) == 200 and tm.colunas_ignoradas == 0, repr(tm.cabecalho))
mantidas = {"uc", "LAT", "LONG", "LAT_GMS", "LONG_GMS", "COORD_X", "COORD_Y", "equipamento_vital", "CLASSE", "ENDERECO",
            "TRANSFORMADOR", "CIRCUITO", "CONJUNTO", "BAIRRO", "SITUACAO"}   # conjunto, bairro e situação: para escolher quem será desligado
conferir("aba Postes: só as 15 colunas do modelo; 26 ignoradas", set(tp.cabecalho) == mantidas and tp.colunas_ignoradas == 26,
         f"{tp.cabecalho} ignoradas={tp.colunas_ignoradas}")
wb = openpyxl.load_workbook(io.BytesIO(postes), read_only=True)
todas = list(wb["Postes"].iter_rows(values_only=True))
cab = todas[0]
texto_saida = repr(tp.cabecalho) + repr(tp.linhas)
sensiveis = ["NOME", "NUM_CASA", "MEDIDOR", "NUM_CONTA", "NUM_ROTA", "NUM_LIVRO"]
vazou = [c for c in sensiveis if c in tp.cabecalho]
vazou += [f"{c}={r[cab.index(c)]!r}" for c in ("NOME",) for r in todas[1:] if repr(r[cab.index(c)]) in texto_saida]
conferir("aba Postes: nenhuma coluna nem valor de NOME, NUM_CASA, MEDIDOR, NUM_CONTA... sai da leitura", not vazou, str(vazou[:5]))
certo = leitura.ler_tabela((EXEMPLOS / "planilha_postes_certo.xlsx").read_bytes())
conferir("planilha_postes_certo (só Postes, sem uc): NUMCDC vira a UC", certo.aba == "Postes"
         and certo.cabecalho[certo.campos["uc"]] == "NUMCDC" and len(certo.linhas) == 200, repr(certo.cabecalho))

# ---------- .csv ----------
t = leitura.ler_tabela("uc;latitude;longitude;local\n1;-21,5;-42,6;Sítio São José\n".encode("cp1252"))
conferir(".csv em Windows-1252 (acentos certos)", t.linhas == [(2, ["1", "-21,5", "-42,6", "Sítio São José"])], repr(t.linhas))
t = leitura.ler_tabela("\ufeffuc,latitude,longitude,local\n1,-21.5,-42.6,\"=HYPERLINK(\"\"http://x\"\")\"\n".encode())
conferir(".csv UTF-8 com BOM, separador vírgula, =HYPERLINK como texto", t.cabecalho[0] == "uc"
         and t.linhas[0][1][3] == '=HYPERLINK("http://x")', repr(t.linhas))
conferir(".csv: dataframe para o validar (textos)", list(t.dataframe().columns) == ["uc", "latitude", "longitude", "local"])

# ---------- limites ----------
grande = "uc;latitude;longitude\n" + "".join(f"{i};-21,5;-42,6\n" for i in range(leitura.MAX_LINHAS + 1))
conferir("mais de 50.000 linhas: recusado", recusa(grande.encode()) == "linhas")
tempo = leitura.TEMPO_MAXIMO
leitura.TEMPO_MAXIMO = -1
conferir("tempo máximo de leitura: recusado", recusa(NORMAL) == "tempo")
leitura.TEMPO_MAXIMO = tempo
conferir("nome de aba limpo (sem controle, até 31)", leitura.nome_de_aba("A\x07b\u200bc  d" + "x" * 40) == "Abc d" + "x" * 26,
         repr(leitura.nome_de_aba("A\x07b\u200bc  d" + "x" * 40)))

n = 50_000
linhas = [CAB] + [[f"{i:010d}", -21.5 - i % 900 / 10_000, -42.6 - i % 700 / 10_000, "não", "residencial", f"Rua {i % 300}"]
                  for i in range(n)]
b = xlsx({"Modelo": linhas})
inicio = time.perf_counter()
t = leitura.ler_tabela(b)
seg = time.perf_counter() - inicio
conferir(f".xlsx de 50.000 linhas lido em {seg:.1f} s ({len(b) / 1e6:.1f} MB)", len(t.linhas) == n and seg < leitura.TEMPO_MAXIMO)

for st, nome, extra in resultados:
    print(st, nome, f"-> {extra}" if st != "OK  " else "")
falhas = sum(1 for r in resultados if r[0] != "OK  ")
print(f"\n{len(resultados) - falhas} de {len(resultados)} verificações OK")
sys.exit(1 if falhas else 0)
