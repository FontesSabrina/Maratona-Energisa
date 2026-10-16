"""Teste da escolha de quem será desligado pela API (nucleo/selecao.py, /api/planilha/escolha), com o banco local.

Num esquema próprio do Postgres (apagado no começo). Confere:
- a base da Energisa inteira (se estiver em data/raw/energisa/, fora do git) e a réplica do git: lidas, inativas fora,
  nenhum número de UC nem nome de titular na resposta;
- o que a tela recebe (base) basta para o contador ao vivo: a conta feita só com ele é a mesma do servidor;
- seleção por área desenhada, por filtro (circuito, bairro), por comunidade e a planilha inteira (inativas fora;
  com incluir_inativas, dentro);
- equipamento vital dentro da escolha bloqueia; fora dela, não; sem coordenada, com o mapa, bloqueia;
- a mesma escolha (valores em outra ordem) dá o mesmo resultado, e o aviso gerado é o que foi conferido;
- o pacote "base" não gera aviso; critério inválido 422; Consulta 403;
- histórico, comprovante e auditoria guardam o critério (com o nome de quem fez), nunca a lista de UCs.
Uso (na raiz, com o banco local): uv run python -W ignore testes/api/teste_escolha.py
"""

import io
import json
import os
import sys
from pathlib import Path

os.environ["FAROL_ESQUEMA"] = "teste_escolha"
os.environ.setdefault("FAROL_EMAIL", "0")

import openpyxl  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from comunidades.api import app as app_mod  # noqa: E402
from comunidades.api.app import app  # noqa: E402
from comunidades.banco.conexao import conectar  # noqa: E402
from comunidades.nucleo import area as area_  # noqa: E402
from comunidades.semente import semear  # noqa: E402

RAIZ = Path(__file__).resolve().parents[2]
REPLICA = RAIZ / "testes" / "planilha" / "exemplos" / "demo" / "6_base_clientes.xlsx"
ENERGISA = RAIZ / "data" / "raw" / "energisa" / "BASE_CLIENTES.xlsx"
FORM = {"data": "20/11", "horario": "8h às 14h"}
resultados = []


def conferir(nome, ok, extra=""):
    resultados.append(("OK  " if ok else "FALHA", nome, extra))


with conectar() as c:
    c.execute("DROP SCHEMA IF EXISTS teste_escolha CASCADE")
temps = {e: s for _, e, s in semear()}


def entrar(email):
    cli = TestClient(app, base_url="http://farol.teste")
    cli.post("/api/entrar", json={"email": email, "senha": temps[email]})
    cli.post("/api/senha", json={"atual": temps[email], "nova": "Lanterna-do-Farol-2026"})
    return cli


op, con = entrar("operador@emg.example"), entrar("consulta@emg.example")


def importar(conteudo: bytes, nome="base.xlsx"):
    app_mod._envios.clear()
    return op.post("/api/planilha", files={"arquivo": (nome, conteudo)}, data=FORM)


def escolher(j, criterio, cli=None):
    return (cli or op).post("/api/planilha/escolha", json={"pacote": j["pacote_base"], "criterio": criterio, **FORM})


def gerar(j):
    return op.post("/api/avisos", json={"pacote": j["pacote"], **FORM, "conferi": True})


def valores(conteudo: bytes, coluna: str) -> list:
    ws = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True).active
    linhas = list(ws.iter_rows(values_only=True))
    i = list(linhas[0]).index(coluna)
    return [r[i] for r in linhas[1:]]


def esperado(base: dict, criterio: dict) -> int:
    """A conta da tela (o contador ao vivo) feita só com o que /api/planilha devolveu em "base"."""
    ln, n = base["linhas"], base["lidas"]
    marcados = {f: {base["valores"][f].index(v) for v in criterio.get(f, [])} for f in ("circuito", "conjunto", "transformador", "bairro")}
    escolhidas_com = {i for i, c in enumerate(base["comunidades"])
                      if (base["municipios"][c["municipio"]]["codigo_ibge"], c["nome"]) in {(x["municipio"], x["nome"]) for x in criterio.get("comunidades", [])}}
    polis = criterio.get("poligonos", [])
    mapa = bool(polis or criterio.get("comunidades"))
    total = 0
    for i in range(n):
        if ln["inativa"][i] and not criterio.get("incluir_inativas"):
            continue
        if criterio["modo"] == "planilha":
            total += 1
            continue
        if any(ln[f][i] not in v for f, v in marcados.items() if v):
            continue
        if mapa:
            tem = ln["lat"][i] is not None
            dentro = tem and any(area_.no_poligono(ln["lat"][i], ln["lon"][i], p) for p in polis)
            if not (dentro or ln["com"][i] in escolhidas_com or ln["vital_sem_posicao"][i]):
                continue
        total += 1
    return total


def caixa_ativas(base):
    ln = base["linhas"]
    pts = sorted((ln["lat"][i], ln["lon"][i]) for i in range(base["lidas"]) if ln["lat"][i] is not None and not ln["inativa"][i])
    lats, lons = sorted(p[0] for p in pts), sorted(p[1] for p in pts)
    a, b = lons[len(lons) // 4], lons[3 * len(lons) // 4]
    c, d = lats[len(lats) // 4], lats[3 * len(lats) // 4]
    return [[[a, c], [b, c], [b, d], [a, d], [a, c]]]


# ---------- a base da Energisa inteira e a réplica ----------
for nome, arquivo, lidas, inativas in [("réplica", REPLICA, 400, None)] + ([("Energisa", ENERGISA, 200, 104)] if ENERGISA.exists() else []):
    conteudo = arquivo.read_bytes()
    r = importar(conteudo, arquivo.name)
    j = r.json() if r.status_code == 200 else {}
    conta = j.get("escolha", {}).get("conta", {})
    inat = inativas if inativas is not None else sum(v == "INATIVO" for v in valores(conteudo, "SITUACAO"))
    conferir(f"{nome} inteira: {lidas} lidas, {inat} inativas fora, {lidas - inat} escolhidas (a planilha inteira)",
             r.status_code == 200 and conta == {"lidas": lidas, "inativas_fora": inat, "nao_escolhidas": 0, "escolhidas": lidas - inat,
                                                "ok": lidas - inat, "corrigidas": 0, "rejeitadas": 0}
             and j["base"]["lidas"] == lidas and j["municipios"], f"{r.status_code} {conta} {r.text[:200]}")
    texto = r.text
    vazou = [str(v) for col in ("NUMCDC", "NOME", "MEDIDOR", "NUM_CONTA") for v in valores(conteudo, col) if str(v) in texto]
    conferir(f"{nome}: nenhum número de UC, titular, medidor ou conta na resposta", not vazou, f"{len(vazou)} valor(es)")
    conferir(f"{nome}: a base para a tela traz os filtros (circuito, conjunto, transformador, bairro e situação)",
             j["base"]["colunas"] == ["circuito", "conjunto", "transformador", "bairro", "situacao"]
             and all(len(j["base"]["linhas"][k]) == lidas for k in j["base"]["linhas"]), str(j["base"]["colunas"]))
    if nome == "Energisa":
        conferir("Energisa: o aviso da planilha inteira sai (96 UCs ativas)", gerar(j).status_code == 200)

conteudo = REPLICA.read_bytes()
j = importar(conteudo).json()
base = j["base"]

# ---------- por área, por filtro, por comunidade ----------
area = caixa_ativas(base)
criterio = {"modo": "escolha", "poligonos": [area]}
r = escolher(j, criterio)
ja = r.json() if r.status_code == 200 else {}
n = esperado(base, criterio)
conferir(f"área desenhada: {n} UCs, a mesma conta da tela (só com a base)", r.status_code == 200 and ja["escolha"]["conta"]["escolhidas"] == n
         and 0 < n < base["lidas"] and ja["conferencia"]["numeros"]["clientes"] <= n, f"{r.status_code} {r.text[:200]}")
conferir("área: a conta fecha (lidas = inativas fora + não escolhidas + escolhidas = ...)",
         (lambda c: c["lidas"] == c["inativas_fora"] + c["nao_escolhidas"] + c["escolhidas"] and c["escolhidas"] == c["ok"] + c["corrigidas"] + c["rejeitadas"])(ja["escolha"]["conta"]))
conferir("área: o texto do critério ('área desenhada (1 polígono, 4 vértices), N UCs, Z comunidades')",
         ja["escolha"]["texto"] == f"área desenhada (1 polígono, 4 vértices), {n} UCs, {ja['conferencia']['numeros']['comunidades']} comunidades",
         ja["escolha"]["texto"])
circ = base["valores"]["circuito"][0]
criterio_c = {"modo": "escolha", "circuito": [circ]}
rc = escolher(j, criterio_c)
conferir(f"filtro circuito {circ}: a mesma conta da tela", rc.status_code == 200 and rc.json()["escolha"]["conta"]["escolhidas"] == esperado(base, criterio_c),
         rc.text[:200])
bairro = max(base["valores"]["bairro"], key=lambda b: base["linhas"]["bairro"].count(base["valores"]["bairro"].index(b)))
criterio_cb = {"modo": "escolha", "circuito": [circ], "bairro": [bairro]}
rcb = escolher(j, criterio_cb)
conferir("circuito e bairro (os dois): a mesma conta da tela", rcb.status_code == 200
         and rcb.json()["escolha"]["conta"]["escolhidas"] == esperado(base, criterio_cb) < rc.json()["escolha"]["conta"]["escolhidas"])
com = base["comunidades"][0]
criterio_com = {"modo": "escolha", "comunidades": [{"municipio": base["municipios"][com["municipio"]]["codigo_ibge"], "nome": com["nome"]}]}
rk = escolher(j, criterio_com)
conferir(f"comunidade clicada ({com['nome']}): a mesma conta da tela e só ela no aviso", rk.status_code == 200
         and rk.json()["escolha"]["conta"]["escolhidas"] == esperado(base, criterio_com) > 0
         and com["nome"] in [c["nome"] for m in rk.json()["municipios"] for c in m["comunidades"]], rk.text[:200])
ri = escolher(j, {"modo": "planilha", "incluir_inativas": True})
conferir("planilha inteira com as inativas: todas as 400", ri.status_code == 200 and ri.json()["escolha"]["conta"]["escolhidas"] == 400)
conferir("inativas fora contadas por bairro e por circuito", sum(j["escolha"]["inativas"]["bairro"].values()) == j["escolha"]["conta"]["inativas_fora"]
         == sum(j["escolha"]["inativas"]["circuito"].values()))

# ---------- a mesma escolha, o mesmo resultado; o aviso é o que foi conferido ----------
outra = {"modo": "escolha", "bairro": [bairro, bairro], "circuito": [circ], "poligonos": [[[[x + 1e-9, y] for x, y in area[0]]]]}
mesma = {"modo": "escolha", "circuito": [circ], "bairro": [bairro], "poligonos": [area]}
r1, r2 = escolher(j, outra).json(), escolher(j, mesma).json()
chave = lambda x: (x["escolha"]["sha256"], x["escolha"]["conta"], x["conferencia"]["numeros"],
                   [(m["nome"], m["ucs"], m["aviso"], [c["nome"] for c in m["comunidades"]]) for m in x["municipios"]])
conferir("a mesma escolha (outra ordem, repetidos) dá o mesmo resultado", chave(r1) == chave(r2))
g = gerar(ja)
gj = g.json() if g.status_code == 200 else {}
conferir("o aviso gerado é o que foi conferido (mesmos municípios, comunidades e aviso)", g.status_code == 200 and
         [(m["nome"], m["ucs"], m["aviso"]) for m in gj["municipios"]] == [(m["nome"], m["ucs"], m["aviso"]) for m in ja["municipios"]], g.text[:200])

# ---------- histórico, comprovante e auditoria: o critério, nunca as UCs ----------
h = op.get(f"/api/historico/{gj.get('id')}").json()
conferir("histórico guarda o critério (texto, sha256 e a conta)", h.get("resumo", {}).get("escolha", {}).get("texto") == ja["escolha"]["texto"]
         and h["resumo"]["escolha"]["sha256"] == ja["escolha"]["sha256"] and "poligonos" in h["resumo"]["escolha"], json.dumps(h.get("resumo"))[:300])
cp = op.get(f"/api/historico/{gj.get('id')}/comprovante").json()
conferir("comprovante traz o critério", cp.get("escolha", {}).get("texto") == ja["escolha"]["texto"])
pdf = op.get(f"/api/historico/{gj.get('id')}/comprovante.pdf")
conferir("comprovante em PDF sai", pdf.status_code == 200 and pdf.content[:4] == b"%PDF")
with conectar() as c:
    regs = c.execute("SELECT acao, detalhe, usuario_nome, perfil FROM registro ORDER BY id").fetchall()
esc = [x for x in regs if x[0] == "escolheu quem será desligado"]
conferir("auditoria: 'escolheu quem será desligado' com o critério e o nome de quem fez",
         len(esc) >= 1 and ja["escolha"]["texto"] in esc[-1][1] and esc[-1][2] and esc[-1][3] == "operador", str(esc[-1:])[:300])
conferir("auditoria: as conferências de escolha também ficam registradas", sum(x[0] == "conferiu a escolha de quem será desligado" for x in regs) >= 5)
ucs = {str(v) for v in valores(conteudo, "NUMCDC")}
conferir("auditoria e histórico sem nenhum número de UC", not any(u in x[1] for x in regs for u in ucs)
         and not any(u in json.dumps(h) for u in ucs))

# ---------- equipamento vital: vale sobre as escolhidas ----------
wb = openpyxl.load_workbook(io.BytesIO(conteudo))
ws = wb.active
cab = [c.value for c in ws[1]]
col = {n: cab.index(n) + 1 for n in cab}
ws.cell(1, len(cab) + 1, "equipamento_vital")
circuitos = base["valores"]["circuito"]
a_linha = next(i for i in range(2, ws.max_row + 1) if ws.cell(i, col["CIRCUITO"]).value == circuitos[0])
b_linha = next(i for i in range(2, ws.max_row + 1) if ws.cell(i, col["CIRCUITO"]).value == circuitos[1])
for linha in (a_linha, b_linha):
    ws.cell(linha, len(cab) + 1, "sim")
    ws.cell(linha, col["SITUACAO"], "ATIVO")
for k in ("LAT", "LONG", "LAT_GMS", "LONG_GMS", "COORD_X", "COORD_Y"):   # a do primeiro circuito sem coordenada: rejeitada
    ws.cell(a_linha, col[k]).value = None   # (ws.cell(l, c, None) não apaga: o openpyxl só grava valor que não é None)
saida = io.BytesIO()
wb.save(saida)
jv = importar(saida.getvalue(), "vital.xlsx").json()
r = escolher(jv, {"modo": "escolha", "circuito": [circuitos[0]]})
cv = r.json()["conferencia"]
conferir("vital rejeitado DENTRO da escolha bloqueia (e o aviso não sai)", [x["linha"] for x in cv["vital_bloqueio"]] == [a_linha]
         and gerar(r.json()).status_code == 422, str(cv["vital_bloqueio"]))
r = escolher(jv, {"modo": "escolha", "circuito": [circuitos[1]]})
conferir("vital FORA da escolha não bloqueia (o aviso sai)", not r.json()["conferencia"]["vital_bloqueio"] and gerar(r.json()).status_code == 200,
         str(r.json()["conferencia"]["vital_bloqueio"]))
r = escolher(jv, {"modo": "escolha", "poligonos": [caixa_ativas(jv["base"])]})
conferir("área: vital sem coordenada entra (não dá para saber se está na área) e bloqueia",
         a_linha in [x["linha"] for x in r.json()["conferencia"]["vital_bloqueio"]])
conferir("a planilha inteira com o vital rejeitado bloqueia (como antes)", bool(jv["conferencia"]["vital_bloqueio"]))

# ---------- o que não pode ----------
r = op.post("/api/avisos", json={"pacote": j["pacote_base"], **FORM, "conferi": True})
conferir("o pacote 'base' não gera aviso (422: escolha primeiro)", r.status_code == 422 and "Escolha quem será desligado" in r.json()["detail"], r.text[:200])
r = op.post("/api/planilha/escolha", json={"pacote": ja["pacote"], "criterio": {"modo": "planilha"}, **FORM})
conferir("o pacote da escolha não serve como base (422)", r.status_code == 422, r.text[:200])
for nome_, crit in [("escolha sem nada", {"modo": "escolha"}), ("campo a mais", {"modo": "planilha", "ucs": ["1"]}),
                    ("anel aberto", {"modo": "escolha", "poligonos": [[[[0, 0], [1, 0], [1, 1]]]]})]:
    r = escolher(j, crit)
    conferir(f"critério inválido: 422 ({nome_})", r.status_code == 422, r.text[:160])
conferir("Consulta não escolhe (403)", escolher(j, {"modo": "planilha"}, con).status_code == 403)
pacote_alterado = j["pacote_base"][:-3] + ("000" if j["pacote_base"][-3:] != "000" else "111")
r = op.post("/api/planilha/escolha", json={"pacote": pacote_alterado, "criterio": {"modo": "planilha"}, **FORM})
conferir("pacote base alterado: 422", r.status_code == 422 and "alterado" in r.json()["detail"])

for st_, nome, extra in resultados:
    print(st_, nome, f"-> {extra}" if st_ != "OK  " and extra else "")
falhas = sum(1 for r in resultados if r[0] != "OK  ")
print(f"\n{len(resultados) - falhas} de {len(resultados)} verificações OK")
sys.exit(1 if falhas else 0)
