"""Teste da base de clientes no tradutor (sem banco e sem API): a planilha que a Energisa mandou é a BASE DE
CLIENTES (colunas EMPRESA, CIRCUITO, CONJUNTO, TRANSFORMADOR, NUMCDC, BAIRRO, SITUACAO, CLASSE, LAT, LONG,
LAT_GMS, LONG_GMS, COORD_X, COORD_Y...), não a lista de um desligamento.

Confere, na réplica do git (testes/planilha/exemplos/demo/6_base_clientes.xlsx, UCs fictícias) e, se existir,
na planilha da Energisa (data/raw/energisa/BASE_CLIENTES.xlsx, fora do git):
- passa no portão, no tradutor e na dupla conferência sem erro (conta fechada, todas as linhas ok);
- sinônimos: NUMCDC é a UC; CIRCUITO, CONJUNTO, TRANSFORMADOR, BAIRRO e SITUACAO são lidos; do ENDERECO
  só a rua; NOME, ENDERECO, NUM_CONTA, NUM_CASA e MEDIDOR descartados (nenhum valor deles no JSON);
- as três fontes de coordenada (decimal, GMS e UTM) dão os mesmos pontos;
- SITUACAO = INATIVO é reconhecida (planilha.inativa).
Uso (na raiz): uv run python -W ignore testes/planilha/teste_base_clientes.py
"""

import io
import math
import sys
from pathlib import Path

import openpyxl

from comunidades.api import conferencia, leitura
from comunidades.nucleo import esquema_planilha as E
from comunidades.nucleo import planilha as P

RAIZ = Path(__file__).resolve().parents[2]
REPLICA = RAIZ / "testes" / "planilha" / "exemplos" / "demo" / "6_base_clientes.xlsx"
ENERGISA = RAIZ / "data" / "raw" / "energisa" / "BASE_CLIENTES.xlsx"
DESCARTADAS = ("NOME", "ENDERECO", "NUM_CONTA", "NUM_CASA", "MEDIDOR")

resultados = []


def conferir(nome, ok, extra=""):
    resultados.append(("OK  " if ok else "FALHA", nome, extra))


def traduzir(conteudo: bytes):
    t = leitura.ler_tabela(conteudo)
    m = P.montar(t.campos, t.cabecalho, t.linhas, P.Parametros(dentro=None, epsg_utm=31983),
                 arquivo={"nome": "base.xlsx", "sha256": "0" * 64, "formato": t.formato, "aba": t.aba,
                          "colunas_ignoradas": t.colunas_ignoradas},
                 empresa="teste", gerado_em="2026-10-16T10:00", gerado_por="teste")
    conferencia.conferir(conteudo, m)   # levanta Divergencia se a rotina independente não bater
    return t, m


def sem_colunas(conteudo: bytes, tirar: set[str]) -> bytes:
    wb = openpyxl.load_workbook(io.BytesIO(conteudo))
    ws = wb.active
    cab = [c.value for c in ws[1]]
    for nome in sorted(tirar, key=lambda n: -cab.index(n)):
        ws.delete_cols(cab.index(nome) + 1)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def valores(conteudo: bytes, coluna: str) -> list:
    ws = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True).active
    linhas = list(ws.iter_rows(values_only=True))
    i = list(linhas[0]).index(coluna)
    return [r[i] for r in linhas[1:]]


def metros(a, b) -> float:
    return math.hypot((a[0] - b[0]) * 110_574, (a[1] - b[1]) * 111_320 * math.cos(math.radians(a[0])))


def base(nome: str, conteudo: bytes, linhas: int, inativas: int, tolerancia_m: float) -> None:
    try:
        t, m = traduzir(conteudo)
    except Exception as e:  # noqa: BLE001
        conferir(f"{nome}: passa no portão, no tradutor e na dupla conferência", False, type(e).__name__)
        return
    r = m.resumo
    conferir(f"{nome}: passa no portão, no tradutor e na dupla conferência ({r.total} linhas)",
             r.total == linhas and r.ok == linhas and r.rejeitadas == 0, r.model_dump_json())
    conferir(f"{nome}: aba Base_Ficticia escolhida sozinha", t.aba == "Base_Ficticia", t.aba)
    esperados = {"uc", "latitude", "longitude", "lat_gms", "lon_gms", "coord_x", "coord_y", "circuito", "conjunto",
                 "transformador", "bairro", "situacao", "classe", "local"}
    conferir(f"{nome}: colunas reconhecidas (NUMCDC, CIRCUITO, CONJUNTO, TRANSFORMADOR, BAIRRO, SITUACAO...)",
             set(m.colunas) == esperados, ", ".join(m.colunas))
    conferir(f"{nome}: o local vem do ENDERECO (a coluna LOCAL é URBANO/RURAL)", t.cabecalho[t.campos["local"]] == "ENDERECO")
    texto = E.texto(m)
    for coluna in ("NOME", "MEDIDOR", "NUM_CONTA"):
        vistos = {str(v) for v in valores(conteudo, coluna) if v not in (None, "")}
        vazou = [v for v in vistos if v in texto]
        conferir(f"{nome}: nenhum valor de {coluna} no JSON", not vazou, f"{len(vazou)} valor(es)")
    # do ENDERECO só a rua (so_rua), e o local de cada ponto é exatamente ela: o NUM_CASA não tem por onde entrar
    # (na base da Energisa o endereço já vem só com a via, como "Rua 17": o número é o nome da via)
    ruas = [P.limpar(P.so_rua(str(v)))[:80] or None for v in valores(conteudo, "ENDERECO")]
    conferir(f"{nome}: ENDERECO e NUM_CASA descartados (o local de cada linha é só a rua do endereço)",
             [p.local for p in m.pontos] == ruas and all(set(p.model_dump()) == set(E.Ponto.model_fields) for p in m.pontos))
    conferir(f"{nome}: {inativas} UCs inativas reconhecidas",
             sum(P.inativa(p.situacao) for p in m.pontos) == inativas == sum(v == "INATIVO" for v in valores(conteudo, "SITUACAO")))
    decimal = [(p.lat, p.lon) for p in m.pontos]
    for fonte, tirar in (("GMS", {"LAT", "LONG", "COORD_X", "COORD_Y"}), ("UTM", {"LAT", "LONG", "LAT_GMS", "LONG_GMS"})):
        _, mf = traduzir(sem_colunas(conteudo, tirar))
        pior = max(metros(a, (p.lat, p.lon)) for a, p in zip(decimal, mf.pontos))
        conferir(f"{nome}: só {fonte} dá os mesmos pontos (pior diferença {pior:.1f} m, limite {tolerancia_m:g} m)",
                 mf.resumo.ok == linhas and pior <= tolerancia_m)


# ---------- sinônimos das colunas novas ----------
for cab, campo in ((["Situação"], "situacao"), (["SITUACAO_UC"], "situacao"), (["status_uc"], "situacao"), (["Bairro"], "bairro"),
                   (["NOME_BAIRRO"], "bairro"), (["Conjunto"], "conjunto"), (["COD_CONJUNTO"], "conjunto")):
    conferir(f"sinônimo: {cab[0]} -> {campo}", campo in P.colunas(["uc", "lat", "long"] + cab))
conferir("inativa: INATIVO, inativa, Desligado; ATIVO, vazio e outro valor ficam ativos",
         all(P.inativa(v) for v in ("INATIVO", "inativa", "Desligado")) and not any(P.inativa(v) for v in ("ATIVO", None, "", "LIGADO")))

base("réplica (git)", REPLICA.read_bytes(), 400, sum(v == "INATIVO" for v in valores(REPLICA.read_bytes(), "SITUACAO")), 0.5)
if ENERGISA.exists():
    base("Energisa", ENERGISA.read_bytes(), 200, 104, 1)
else:
    print(f"(planilha da Energisa ausente em {ENERGISA}: só a réplica)")

for st_, nome, extra in resultados:
    print(st_, nome, f"-> {extra}" if st_ != "OK  " and extra else "")
falhas = sum(1 for r in resultados if r[0] != "OK  ")
print(f"\n{len(resultados) - falhas} de {len(resultados)} verificações OK")
sys.exit(1 if falhas else 0)
