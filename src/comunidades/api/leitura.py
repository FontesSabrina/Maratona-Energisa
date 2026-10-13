"""Portão e leitura segura do arquivo enviado (.xlsx ou .csv), em memória. Nada é gravado em disco.

Portão (antes de abrir o arquivo):
- o tipo vem dos bytes iniciais, nunca do nome: .xlsx = zip ("PK"); .csv = texto. Recusados: .xls
  antigo e .xlsx protegido por senha (contêiner OLE, D0 CF 11 E0), executável ("MZ"), outros
  binários, bytes nulos no .csv, .xlsm (vbaProject.bin ou "macroEnabled"), .xlsb, zip cifrado;
- tamanho até 10 MB; no zip: até 50 MB descompactado, razão de compressão até 100, até 200 arquivos
  internos, nenhum caminho com ".." ou começando com "/" ou "\\". O zipfile do Python não descompacta
  além do tamanho declarado no índice (passou disso, o arquivo é recusado como corrompido).
Leitura:
- openpyxl em modo só leitura (read_only), com os valores guardados (data_only) e o XML pelo
  defusedxml; nada de eval, pickle ou pandas.read_excel. Uma segunda passada sem data_only marca as
  fórmulas sem valor guardado (FORMULA_SEM_VALOR), para a linha ser rejeitada depois;
- só abas visíveis; a aba é escolhida sozinha quando só uma tem as colunas obrigatórias; com mais de
  uma, EscolhaDeAba devolve a lista para a pessoa escolher;
- até 50.000 linhas e 60 colunas, e um tempo máximo de leitura;
- minimização já aqui: só as colunas do modelo (nucleo/planilha.MANTIDAS) saem da leitura; as outras
  (nome, número da casa, medidor, conta...) são contadas e descartadas sem serem copiadas;
- .csv em UTF-8 (com ou sem BOM), senão Windows-1252; separador ; ou ,.
As mensagens de recusa são fixas (MENSAGENS) e nunca trazem conteúdo do arquivo.
"""

import csv
import io
import time
import unicodedata
import zipfile
from dataclasses import dataclass, field

import openpyxl
import pandas as pd

from ..nucleo import planilha as regras
from ..nucleo.motivos import FORMULA_SEM_VALOR
from ..nucleo.planilha import PlanilhaInvalida

MAX_BYTES = 10 * 1024 * 1024
MAX_DESCOMPACTADO = 50 * 1024 * 1024   # soma do conteúdo descompactado (50.000 linhas usam ~18 MB)
MAX_ENTRADAS = 200                     # arquivos dentro do .xlsx (um típico tem uns 10 a 30)
MAX_RAZAO = 100                        # descompactado / compactado de cada arquivo (XML de planilha: ~5 a 15)
MAX_LINHAS = regras.MAX_LINHAS
MAX_COLUNAS = 60
TEMPO_MAXIMO = 60.0                    # segundos de leitura (50.000 linhas levam poucos segundos)
MAX_NOME_ABA = 31                      # o Excel não deixa passar disso

ZIP, OLE, EXE = b"PK\x03\x04", bytes.fromhex("D0CF11E0A1B11AE1"), b"MZ"
BINARIOS = (b"%PDF", b"\x89PNG", b"GIF8", b"\xff\xd8\xff", b"\x1f\x8b", b"7z\xbc\xaf", b"Rar!", b"\x7fELF")

MENSAGENS = {
    "vazio": "O arquivo está vazio.",
    "grande": "O arquivo passa de 10 MB.",
    "ole": "Arquivo .xls antigo ou planilha protegida por senha: salve como .xlsx sem senha (ou .csv) e envie de novo.",
    "nao_planilha": "O arquivo não é uma planilha .xlsx ou .csv.",
    "zip": "Não foi possível ler o .xlsx (arquivo corrompido).",
    "cifrado": "O .xlsx está protegido por senha: salve sem senha e envie de novo.",
    "macro": "Planilha com macro (.xlsm) não é aceita: salve como .xlsx sem macros e envie de novo.",
    "xlsb": "Planilha binária (.xlsb) não é aceita: salve como .xlsx e envie de novo.",
    "entradas": f"O .xlsx tem arquivos internos demais (mais de {MAX_ENTRADAS}).",
    "descompactado": "O .xlsx descompactado passa de 50 MB.",
    "razao": "O .xlsx tem compactação fora do normal (possível arquivo malicioso).",
    "caminho": "O .xlsx tem caminhos internos inválidos (possível arquivo malicioso).",
    "linhas": f"A planilha passa de {MAX_LINHAS:,} linhas.".replace(",", "."),
    "colunas": f"A planilha passa de {MAX_COLUNAS} colunas.",
    "tempo": "A leitura passou do tempo máximo. Divida a planilha em partes menores e envie de novo.",
    "nul": "O .csv tem caracteres inválidos (bytes nulos): não parece um arquivo de texto.",
    "codificacao": "Não foi possível ler o .csv (codificação desconhecida: use UTF-8 ou a do Excel).",
    "csv": "Não foi possível ler o .csv (formato inesperado).",
    "sem_aba": "Nenhuma aba visível tem as colunas obrigatórias (uc e coordenadas).",
    "sem_colunas": "Faltam colunas obrigatórias: uc e as coordenadas (latitude e longitude). "
                   "O modelo tem: uc, latitude, longitude, equipamento_vital, classe e local.",
    "aba_invalida": "A aba escolhida não existe ou não tem as colunas obrigatórias.",
}


class Recusa(PlanilhaInvalida):
    """O arquivo inteiro foi recusado no portão ou na leitura (mensagem fixa de MENSAGENS)."""

    def __init__(self, codigo: str):
        super().__init__(MENSAGENS[codigo])
        self.codigo = codigo


class EscolhaDeAba(PlanilhaInvalida):
    """Mais de uma aba visível tem as colunas obrigatórias: a pessoa escolhe qual usar."""

    def __init__(self, abas: list[str]):
        super().__init__("Mais de uma aba tem as colunas obrigatórias. Escolha qual usar: " + ", ".join(abas) + ".")
        self.abas = abas


@dataclass
class Tabela:
    """O que sai da leitura: só as colunas do modelo, com os valores como estavam nas células."""
    formato: str                         # xlsx | csv
    aba: str | None                      # a aba usada (None no .csv)
    cabecalho: list[str]                 # nomes originais das colunas mantidas
    campos: dict[str, int]               # campo do modelo -> posição em cabecalho
    linhas: list[tuple[int, list]] = field(default_factory=list)   # (número da linha na planilha, valores)
    colunas_ignoradas: int = 0           # colunas descartadas (só a contagem)

    def dataframe(self) -> pd.DataFrame:
        """As linhas como tabela de textos, com os nomes originais (para nucleo/planilha.validar)."""
        texto = lambda v: "" if v is None or v is FORMULA_SEM_VALOR else str(v)
        return pd.DataFrame([[texto(v) for v in valores] for _, valores in self.linhas], columns=self.cabecalho, dtype=str)


# ---------- portão ----------
def conferir_xlsx(conteudo: bytes) -> None:
    """O índice do zip, antes de abrir o .xlsx: tamanho, razão, entradas, caminhos, macro, cifra."""
    try:
        with zipfile.ZipFile(io.BytesIO(conteudo)) as z:
            itens = z.infolist()
            nomes = {i.filename for i in itens}
            tipos = z.read("[Content_Types].xml") if "[Content_Types].xml" in nomes and len(itens) <= MAX_ENTRADAS \
                and z.getinfo("[Content_Types].xml").file_size <= 1024 * 1024 else b""
    except (zipfile.BadZipFile, ValueError, OSError, RuntimeError, NotImplementedError):
        raise Recusa("zip") from None
    if len(itens) > MAX_ENTRADAS:
        raise Recusa("entradas")
    if any(i.filename.startswith(("/", "\\")) or ".." in i.filename.replace("\\", "/").split("/")
           or ":" in i.filename for i in itens):
        raise Recusa("caminho")
    if any(i.flag_bits & 0x1 for i in itens):
        raise Recusa("cifrado")
    if sum(i.file_size for i in itens) > MAX_DESCOMPACTADO:
        raise Recusa("descompactado")
    if any(i.file_size > MAX_RAZAO * max(i.compress_size, 1) for i in itens):
        raise Recusa("razao")
    if any(n.lower().endswith("vbaproject.bin") for n in nomes) or b"macroenabled" in tipos.lower():
        raise Recusa("macro")
    if "xl/workbook.bin" in nomes:
        raise Recusa("xlsb")
    if "xl/workbook.xml" not in nomes:
        raise Recusa("nao_planilha")


def tipo(conteudo: bytes) -> str:
    """xlsx ou csv, pelos bytes iniciais; o resto é recusado."""
    if not conteudo.strip():
        raise Recusa("vazio")
    if len(conteudo) > MAX_BYTES:
        raise Recusa("grande")
    if conteudo.startswith(ZIP):
        return "xlsx"
    if conteudo.startswith(OLE):
        raise Recusa("ole")
    if conteudo.startswith(EXE) or conteudo.startswith(BINARIOS):
        raise Recusa("nao_planilha")
    if b"\x00" in conteudo:
        raise Recusa("nul")
    return "csv"


# ---------- leitura ----------
class _Relogio:
    def __init__(self):
        self.fim = time.monotonic() + TEMPO_MAXIMO

    def conferir(self) -> None:
        if time.monotonic() > self.fim:
            raise Recusa("tempo")


def _cabecalho(linha) -> list[str]:
    nomes = ["" if v is None else str(v).strip() for v in (linha or ())]
    while nomes and not nomes[-1]:
        nomes.pop()
    if len(nomes) > MAX_COLUNAS:
        raise Recusa("colunas")
    return nomes


def nome_de_aba(nome: str) -> str:
    """O nome da aba como pode aparecer na tela: sem caracteres de controle e invisíveis, até 31 caracteres."""
    limpo = "".join(c for c in str(nome) if unicodedata.category(c)[0] != "C")
    return " ".join(limpo.split())[:MAX_NOME_ABA]


def _montar(formato: str, aba, cabecalho: list[str], linhas_brutas, relogio: _Relogio) -> Tabela:
    """Mantém só as colunas do modelo e descarta as linhas totalmente vazias."""
    campos_orig = regras.campos_por_indice(cabecalho)
    indices = sorted(set(campos_orig.values()))
    nova_pos = {i: p for p, i in enumerate(indices)}
    t = Tabela(formato, aba, [cabecalho[i] for i in indices], {c: nova_pos[i] for c, i in campos_orig.items()},
               colunas_ignoradas=sum(1 for n in cabecalho if n) - len(indices))
    for n, (numero, valores) in enumerate(linhas_brutas):
        if n % 1000 == 0:
            relogio.conferir()
        mantidos = [valores[i] if i < len(valores) else None for i in indices]
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in mantidos):
            continue
        if len(t.linhas) >= MAX_LINHAS:
            raise Recusa("linhas")
        t.linhas.append((numero, mantidos))
    return t


def _xlsx(conteudo: bytes, aba: str | None, relogio: _Relogio) -> Tabela:
    conferir_xlsx(conteudo)
    try:
        valores = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
        formulas = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=False)
    except Exception:  # noqa: BLE001 (qualquer falha do openpyxl: arquivo inválido, mensagem fixa)
        raise Recusa("zip") from None
    try:
        servem = []
        for titulo in valores.sheetnames:
            ws = valores[titulo]
            if getattr(ws, "sheet_state", "visible") != "visible" or not hasattr(ws, "iter_rows"):
                continue
            cab = _cabecalho(next(ws.iter_rows(max_row=1, values_only=True), ()))
            if regras.serve(regras.campos_por_indice(cab)):
                servem.append(titulo)
        if aba is not None:
            escolhida = next((t for t in servem if nome_de_aba(t) == aba), None)
            if escolhida is None:
                raise Recusa("aba_invalida")
        elif len(servem) == 1:
            escolhida = servem[0]
        elif not servem:
            raise Recusa("sem_aba")
        else:
            raise EscolhaDeAba([nome_de_aba(t) for t in servem])
        linhas_v = valores[escolhida].iter_rows(values_only=True)
        linhas_f = formulas[escolhida].iter_rows(values_only=True)
        cabecalho = _cabecalho(next(linhas_v, ()))
        next(linhas_f, None)

        def brutas():
            for numero, (v, f) in enumerate(zip(linhas_v, linhas_f), start=2):
                yield numero, [FORMULA_SEM_VALOR if a is None and isinstance(b, str) and b.startswith("=") else a
                               for a, b in zip(v, f)]
        return _montar("xlsx", nome_de_aba(escolhida), cabecalho, brutas(), relogio)
    except (PlanilhaInvalida, MemoryError):
        raise
    except Exception:  # noqa: BLE001
        raise Recusa("zip") from None
    finally:
        valores.close()
        formulas.close()


def _csv(conteudo: bytes, relogio: _Relogio) -> Tabela:
    try:
        texto = conteudo.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            texto = conteudo.decode("cp1252")
        except UnicodeDecodeError:
            raise Recusa("codificacao") from None
    primeira = texto.lstrip("﻿").split("\n", 1)[0]
    sep = ";" if primeira.count(";") >= primeira.count(",") else ","
    try:
        leitor = csv.reader(io.StringIO(texto, newline=""), delimiter=sep)
        cabecalho = _cabecalho(next(leitor, []))
        if not regras.serve(regras.campos_por_indice(cabecalho)):
            raise Recusa("sem_colunas")
        return _montar("csv", None, cabecalho, ((n, linha) for n, linha in enumerate(leitor, start=2)), relogio)
    except csv.Error:
        raise Recusa("csv") from None


csv.field_size_limit(MAX_BYTES)


def ler_tabela(conteudo: bytes, aba: str | None = None) -> Tabela:
    """Portão e leitura: a tabela só com as colunas do modelo, ou Recusa / EscolhaDeAba."""
    relogio = _Relogio()
    return _xlsx(conteudo, aba, relogio) if tipo(conteudo) == "xlsx" else _csv(conteudo, relogio)


def ler(conteudo: bytes, aba: str | None = None) -> pd.DataFrame:
    """A tabela como textos, para nucleo/planilha.validar."""
    return ler_tabela(conteudo, aba).dataframe()
